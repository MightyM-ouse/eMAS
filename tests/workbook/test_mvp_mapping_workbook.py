"""Content, native XLSX and negative checks for the v4.4 MVP baseline."""
from __future__ import annotations

import copy
import hashlib
import sys
import tempfile
import unittest
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "build"))
from emas_mvp_derivation import derive_scenario
from emas_mvp_model import PHASES, SHEETS, columns, source_rows
from emas_mvp_validation import validate_model
from generate_emas_mapping_workbook import build, load_json_field_map, table_name
from validate_emas_mapping_workbook import inspect, table_rows


class WorkbookMvpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="emas-mvp-tests-")
        cls.path = Path(cls.temp.name) / "mapping.xlsx"
        cls.summary = build(output=cls.path)
        cls.wb = load_workbook(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()
        cls.temp.cleanup()

    def test_workbook_contract_and_native_features(self):
        self.assertEqual(self.wb.sheetnames, list(SHEETS))
        self.assertEqual(sum(len(ws.tables) for ws in self.wb), 27)
        for sheet in SHEETS[1:]:
            ws = self.wb[sheet]
            self.assertIn(table_name(sheet), ws.tables)
            self.assertEqual([ws.cell(4, i).value for i in range(1, len(columns()[sheet]) + 1)], columns()[sheet])
            self.assertIsNotNone(ws.tables[table_name(sheet)].autoFilter)
            self.assertIsNotNone(ws.freeze_panes)
        self.assertEqual(self.wb["06_Requirement_Catalogue"].freeze_panes, "E5")
        self.assertGreater(len(self.wb["05_Scenario_Module_Map"].data_validations.dataValidation), 0)
        self.assertIn("VL_Applicability", self.wb.defined_names)
        self.assertIn("SelectedScenarioId", self.wb.defined_names)
        with zipfile.ZipFile(self.path) as archive:
            self.assertFalse(any("vba" in name.lower() for name in archive.namelist()))

    def test_approved_content_counts_and_keys(self):
        scenario = table_rows(self.wb["01_Migration_Scenarios"])
        modules = table_rows(self.wb["04_Assessment_Modules"])
        mapping = table_rows(self.wb["05_Scenario_Module_Map"])
        self.assertEqual(len(scenario), 8)
        self.assertEqual(len(modules), 15)
        self.assertEqual(len(mapping), 360)
        self.assertEqual(len(table_rows(self.wb["02_Scenario_Questionnaire"])), 23)
        self.assertEqual(len(table_rows(self.wb["06_Requirement_Catalogue"])), 55)
        self.assertEqual(len({r["RequirementId"] for r in table_rows(self.wb["06_Requirement_Catalogue"]) }), 55)
        self.assertEqual(len({r["FieldCode"] for r in table_rows(self.wb["07_Fields_Evidence"]) }), 25)
        combinations = Counter((r["ScenarioId"], r["Phase"], r["ModuleId"]) for r in mapping)
        self.assertEqual(len(combinations), 8 * 3 * 15)
        self.assertTrue(all(count == 1 for count in combinations.values()))
        self.assertEqual({r["Applicability"] for r in mapping}, {"Required", "Conditional", "Optional", "NotApplicable"})
        self.assertTrue(all(r["ActivationContextField"] and r["ActivationOperator"] and r["ActivationValue"] for r in mapping if r["Applicability"] == "Conditional"))
        self.assertTrue(all(r["Applicability"] == "NotApplicable" for r in mapping if r["ModuleId"] == "MOD-READINESS" and r["Phase"] != "PreMigration"))
        self.assertTrue(all(r["Applicability"] == "NotApplicable" for r in mapping if r["ModuleId"] == "MOD-RECONCILE" and (r["Phase"] != "PostMigration" or r["ScenarioId"] == "MS-07")))
        hybrid = {r["ModuleId"]: r for r in mapping if r["ScenarioId"] == "MS-05" and r["Phase"] == "PreMigration"}
        self.assertEqual(hybrid["MOD-DB"]["ActivationValue"], "ECTDManagerDatabaseArchive")
        self.assertEqual(hybrid["MOD-ARCHIVE"]["ActivationValue"], "ECTDManagerDatabaseArchive")
        self.assertEqual(hybrid["MOD-DMS"]["ActivationValue"], "DMS")

    def test_derivation_including_unsupported_dms_target(self):
        sql = {"TargetPlatform": "eCTDmanager", "PrimarySourceMechanism": "ECTDManagerDatabaseArchive", "SourceDatabaseType": "SQLServer", "ArchiveAvailable": "Yes"}
        self.assertEqual(derive_scenario(sql)["scenarioId"], "MS-01")
        self.assertEqual(derive_scenario({**sql, "ArchiveAvailable": "No"})["status"], "DerivedWithFollowUp")
        self.assertEqual(derive_scenario({**sql, "SourceDatabaseType": "Access"})["scenarioId"], "MS-02")
        self.assertEqual(derive_scenario({"TargetPlatform": "eCTDmanager", "PrimarySourceMechanism": "DMS"})["scenarioId"], "MS-08")
        excluded = derive_scenario({"TargetPlatform": "DMS", "PrimarySourceMechanism": "DMS"})
        self.assertEqual((excluded["scenarioId"], excluded["status"], excluded["reasonCode"]), ("MS-07", "NeedsReview", "OUTSIDE_SUPPORTED_TARGET"))
        self.assertTrue(excluded["consultantDiscussionRequired"])
        self.assertEqual(derive_scenario({"TargetPlatform": "ToBeDefined", "PrimarySourceMechanism": "Unknown"})["followUpQuestionIds"], ["Q-SCN-021", "Q-SCN-022"])
        hybrid = derive_scenario({"TargetPlatform": "eCTDmanager", "PrimarySourceMechanism": "MultipleSources", "IncludedSourceMechanisms": ["ECTDManagerDatabaseArchive", "DMS"]})
        self.assertEqual(hybrid["scenarioId"], "MS-05")
        incomplete = derive_scenario({"TargetPlatform": "eCTDmanager", "PrimarySourceMechanism": "MultipleSources", "IncludedSourceMechanisms": ["DMS"]})
        self.assertFalse(incomplete["runtimeJsonEligible"])

    def test_statuses_are_independent_and_review_gate_visible(self):
        rows = table_rows(self.wb["06_Requirement_Catalogue"])
        self.assertTrue(all({"RequirementStatus", "ImplementationStatus", "VerificationStatus"} <= set(r) for r in rows))
        self.assertTrue(any(r["RequirementStatus"] == "Approved" and r["ImplementationStatus"] == "Implemented" and r["VerificationStatus"] == "NotTested" for r in rows))
        validation = table_rows(self.wb["27_Validation_Results"])
        self.assertTrue(any(r["ControlCode"] == "SOURCE_OBLIGATION_RECONCILIATION_PENDING" and r["Blocking"] for r in validation))
        self.assertEqual(self.wb["00_Home"]["D6"].value, "Blocked")
        self.assertGreater(len(table_rows(self.wb["24_Final_Config_Master"])), 360)
        self.assertTrue(table_rows(self.wb["26_JSON_Preview"]))
        self.assertFalse(inspect(self.path)["failures"])

    def test_negative_semantic_checks(self):
        rows = source_rows()
        rows["25_JSON_Field_Map"] = load_json_field_map()
        self.assertEqual([i.code for i in validate_model(rows) if i.blocking], ["SOURCE_OBLIGATION_RECONCILIATION_PENDING"])
        broken = copy.deepcopy(rows)
        broken["05_Scenario_Module_Map"][0]["ActivationContextField"] = "MISSING-FIELD"
        broken["05_Scenario_Module_Map"][0]["Applicability"] = "Conditional"
        broken["05_Scenario_Module_Map"][0]["ActivationOperator"] = "Equals"
        broken["05_Scenario_Module_Map"][0]["ActivationValue"] = "Yes"
        broken["05_Scenario_Module_Map"].append(copy.deepcopy(broken["05_Scenario_Module_Map"][0]))
        codes = {i.code for i in validate_model(broken)}
        self.assertTrue({"DUPLICATE_ID", "MAP_360", "ACTIVATION_FIELD"} <= codes)
        broken["25_JSON_Field_Map"][0]["SourceColumn"] = "MISSING-COLUMN"
        self.assertIn("JSON_MAP_SOURCE", {i.code for i in validate_model(broken)})

    def test_deterministic_regeneration(self):
        other = Path(self.temp.name) / "mapping-again.xlsx"
        build(output=other)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).digest(), hashlib.sha256(other.read_bytes()).digest())

    def test_other_scenario_review_views(self):
        for scenario in ("MS-02", "MS-03", "MS-04", "MS-05", "MS-06", "MS-07", "MS-08"):
            path = Path(self.temp.name) / f"{scenario}.xlsx"
            build(selected=scenario, output=path)
            self.assertFalse(inspect(path)["failures"])
            wb = load_workbook(path)
            self.assertEqual(wb["00_Home"]["D5"].value, scenario)
            mappings = [r for r in table_rows(wb["05_Scenario_Module_Map"]) if r["ScenarioId"] == scenario]
            self.assertEqual(len(mappings), 45)
            if scenario == "MS-07":
                readiness = next(r for r in mappings if r["Phase"] == "PreMigration" and r["ModuleId"] == "MOD-READINESS")
                reconcile = next(r for r in mappings if r["Phase"] == "PostMigration" and r["ModuleId"] == "MOD-RECONCILE")
                self.assertEqual(readiness["Applicability"], "Required")
                self.assertEqual(readiness["DefaultMissingEvidenceOutcome"], "Blocked")
                self.assertEqual(reconcile["Applicability"], "NotApplicable")
            wb.close()


if __name__ == "__main__":
    unittest.main()
