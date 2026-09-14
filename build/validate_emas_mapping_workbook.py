#!/usr/bin/env python3
"""Inspect the saved v4.4 MVP workbook, not just the in-memory seed."""
from __future__ import annotations

import argparse
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from emas_mvp_model import PHASES, SHEETS, columns, modules, scenarios
from generate_emas_mapping_workbook import OUTPUT, table_name


def table_rows(ws) -> list[dict]:
    headers = [cell.value for cell in ws[4]]
    result = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if any(v is not None for v in row):
            result.append(dict(zip(headers, row)))
    return result


def inspect(path: Path) -> dict:
    failures: list[str] = []
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=False)
    expected_cols = columns()
    if tuple(wb.sheetnames) != SHEETS:
        failures.append("The 28 required sheets are missing or out of order.")
    for sheet in SHEETS[1:]:
        if sheet not in wb:
            continue
        ws = wb[sheet]
        actual = [ws.cell(4, i).value for i in range(1, len(expected_cols[sheet]) + 1)]
        if actual != expected_cols[sheet]:
            failures.append(f"{sheet}: required columns differ from v4.4.")
        if table_name(sheet) not in ws.tables:
            failures.append(f"{sheet}: named Excel Table missing.")
        elif ws.tables[table_name(sheet)].autoFilter is None:
            failures.append(f"{sheet}: table AutoFilter missing.")
        if ws.freeze_panes is None:
            failures.append(f"{sheet}: header rows are not frozen.")
    scenarios_rows = table_rows(wb["01_Migration_Scenarios"])
    module_rows = table_rows(wb["04_Assessment_Modules"])
    mappings = table_rows(wb["05_Scenario_Module_Map"])
    requirements = table_rows(wb["06_Requirement_Catalogue"])
    if {r["ScenarioId"] for r in scenarios_rows if r["IsActive"]} != {r["ScenarioId"] for r in scenarios()} or len(scenarios_rows) != 8:
        failures.append("Scenario catalogue differs from the approved eight.")
    if {r["ModuleId"] for r in module_rows if r["IsActive"]} != {r["ModuleId"] for r in modules()} or len(module_rows) != 15:
        failures.append("Module catalogue differs from the approved fifteen.")
    map_keys = Counter((r["ScenarioId"], r["Phase"], r["ModuleId"]) for r in mappings if r["IsActive"])
    required_keys = {(s["ScenarioId"], p, m["ModuleId"]) for s in scenarios_rows for p in PHASES for m in module_rows}
    if len(mappings) != 360 or set(map_keys) != required_keys or any(count != 1 for count in map_keys.values()):
        failures.append("Scenario/phase/module map is not the full 360 unique records.")
    if any(r["ModuleId"] == "MOD-READINESS" and r["Phase"] != "PreMigration" and r["Applicability"] != "NotApplicable" for r in mappings):
        failures.append("MOD-READINESS appears outside PreMigration.")
    if any(r["ModuleId"] == "MOD-RECONCILE" and (r["Phase"] != "PostMigration" or r["ScenarioId"] == "MS-07") and r["Applicability"] != "NotApplicable" for r in mappings):
        failures.append("MOD-RECONCILE violates its phase or MS-07 boundary.")
    if len({r["RequirementId"] for r in requirements}) != len(requirements):
        failures.append("RequirementId is duplicated.")
    if len({r["FieldCode"] for r in table_rows(wb["07_Fields_Evidence"])}) != len(table_rows(wb["07_Fields_Evidence"])):
        failures.append("FieldCode is duplicated.")
    derivation = table_rows(wb["03_Scenario_Derivation_Rules"])
    outside = [r for r in derivation if r["DerivationRuleId"] == "SDR-TARGET-UNSUPPORTED"]
    if {r["ExpectedValue"] for r in outside} != {"DMS", "ThirdPartySystem", "OtherEXTEDOProduct"} or any(r["CandidateScenarioId"] != "MS-07" or r["OnMatchStatus"] != "NeedsReview" or r["ReasonCode"] != "OUTSIDE_SUPPORTED_TARGET" for r in outside):
        failures.append("Unsupported target derivation is not MS-07 / NeedsReview.")
    if not wb["05_Scenario_Module_Map"].data_validations.dataValidation or not wb["06_Requirement_Catalogue"].data_validations.dataValidation or "VL_Applicability" not in wb.defined_names:
        failures.append("Controlled-value dropdowns or named list ranges are missing.")
    if wb["06_Requirement_Catalogue"].freeze_panes != "E5":
        failures.append("The first four requirement columns are not frozen.")
    if wb["00_Home"]["D5"].value not in {r["ScenarioId"] for r in scenarios_rows}:
        failures.append("Home scenario selection is invalid.")
    with zipfile.ZipFile(path) as archive:
        if any("vba" in name.lower() for name in archive.namelist()):
            failures.append("The XLSX contains VBA or macro parts.")
    validation = table_rows(wb["27_Validation_Results"])
    if not any(r["ControlCode"] == "SOURCE_OBLIGATION_RECONCILIATION_PENDING" and r["Blocking"] for r in validation):
        failures.append("The draft source-obligation gate is not visible as blocking.")
    selected = wb["00_Home"]["D5"].value
    master = table_rows(wb["24_Final_Config_Master"])
    preview = {r["Section"]: r for r in table_rows(wb["26_JSON_Preview"])}
    selected_master_maps = [r for r in master if r["SourceSheet"] == "05_Scenario_Module_Map" and r["SelectedScenarioId"] == selected and r["SourceRecordId"].startswith("SMM-" + str(selected).replace("-", ""))]
    if len(selected_master_maps) != 45 or preview.get("modules[]", {}).get("ObjectCount") != 45:
        failures.append("Final Config Master and JSON Preview disagree on the selected 45 module objects.")
    if preview.get("questionnaire.questions[]", {}).get("ObjectCount") != len(table_rows(wb["02_Scenario_Questionnaire"])):
        failures.append("JSON Preview question count differs from the maintained questionnaire.")
    field_maps = table_rows(wb["25_JSON_Field_Map"])
    for row in field_maps:
        if row["SourceSheet"] not in wb or row["SourceColumn"] not in expected_cols.get(row["SourceSheet"], ()) or row["SourceTable"] != table_name(row["SourceSheet"]):
            failures.append(f"JSON field map {row['MappingId']} has unresolved source lineage.")
            break
    return {"path": str(path), "sheets": len(wb.sheetnames), "tables": sum(len(ws.tables) for ws in wb),
        "scenarios": len(scenarios_rows), "modules": len(module_rows), "mappings": len(mappings),
        "requirements": len(requirements), "validationResults": len(validation), "failures": failures}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", nargs="?", type=Path, default=OUTPUT)
    args = parser.parse_args()
    result = inspect(args.path)
    print(json.dumps(result, indent=2))
    if result["failures"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
