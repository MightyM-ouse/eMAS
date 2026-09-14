#!/usr/bin/env python3
"""Regenerate the macro-free eMAS Mapping Workbook MVP v0.1.

Run from the repository root: python build/generate_emas_mapping_workbook.py
Install the build-only dependency first: python -m pip install -r build/requirements-mvp-workbook.txt
No Microsoft Excel, VBA, network service, or customer evidence is used.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import tempfile
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

from emas_mvp_model import GENERATED, PHASES, ROOT, SHEETS, SOURCE_ID, columns, source_rows
from emas_mvp_validation import Issue, validate_model

OUTPUT = ROOT / "dist/eMAS_Mapping_Workbook_MVP_v0.1.xlsx"
FIELD_MAP_SOURCE = ROOT / "config/authoring/mvp/json-field-map.json"
NAVY, BLUE, LIGHT, AMBER = "172B46", "245F8D", "EAF2F8", "FFF1CF"
WHITE, TEXT, MUTED = "FFFFFF", "203044", "64748B"
JSON_ROOTS = {
    "01_Migration_Scenarios": "scenario", "02_Scenario_Questionnaire": "questionnaire.questions[]",
    "03_Scenario_Derivation_Rules": "questionnaire.derivationRules[]", "04_Assessment_Modules": "modules[]",
    "05_Scenario_Module_Map": "modules[]", "06_Requirement_Catalogue": "requirements[]",
    "07_Fields_Evidence": "catalogues.fields[]", "08_Regulatory_Profiles": "catalogues.regulatoryProfiles[]",
    "09_Dossier_Sequence_ID": "rules.dossierSequenceIdentification[]", "10_Folder_File_Structure": "rules.folderFileStructure[]",
    "11_Missing_Refs_Integrity": "rules.referenceIntegrity[]", "12_Technical_Observations": "rules.technicalObservations[]",
    "13_Size_Volume_Metrics": "rules.sizeVolumeMetrics[]", "14_Source_DB_Archive_DMS": "sourceMappings[]",
    "15_RAG_Severity": "interpretation.ragSeverityRules[]", "16_Confidence": "interpretation.confidenceRules[]",
    "17_Effort_Drivers": "interpretation.effortDrivers[]", "18_Findings": "findings[]",
    "19_Recommendations_Actions": "recommendations[]", "20_PreMigration_Readiness": "phaseRules.preMigration[]",
    "21_PostMigration_Reconciliation": "phaseRules.postMigration[]", "22_Value_Lists": "valueLists",
    "23_Source_References": "sources[]",
}
CONTROLLED_COLUMNS = {
    "Phase": "Phase", "Applicability": "Applicability", "AssessmentDepth": "AssessmentDepth",
    "RulePurpose": "RulePurpose", "OnMatchStatus": "OnMatchStatus", "MissingInputAction": "MissingInputAction",
    "DefaultMissingEvidenceOutcome": "DefaultMissingEvidenceOutcome", "PhaseOutcomeImpact": "PhaseOutcomeImpact",
    "BaselineContribution": "BaselineContribution", "ReconciliationRole": "ReconciliationRole",
    "ModuleLayer": "ModuleLayer", "ExecutionMode": "ExecutionMode",
    "RequirementDomain": "RequirementDomain", "RequirementType": "RequirementType",
    "ObligationLevel": "ObligationLevel", "OwningComponent": "OwningComponent",
    "LifecycleStage": "LifecycleStage", "PhaseScope": "PhaseScope", "ApplicabilityBasis": "ApplicabilityBasis",
    "MissingEvidenceBehavior": "MissingEvidenceBehavior", "ImplementationDisposition": "ImplementationDisposition",
    "VerificationMethod": "VerificationMethod", "RequirementStatus": "RequirementStatus",
    "ImplementationStatus": "ImplementationStatus", "VerificationStatus": "VerificationStatus",
    "DataType": "DataType", "EvidenceSourceType": "EvidenceSourceType",
    "MissingValueMeaning": "MissingValueMeaning", "Aggregation": "Aggregation", "Unit": "Unit",
    "AnswerType": "QuestionAnswerType", "AnswerOwner": "AnswerOwner",
    "MissingAnswerImpact": "MissingAnswerImpact", "PreSalesDetailLevel": "PreSalesDetailLevel",
    "Operator": "Operator", "ActivationOperator": "Operator", "AllowedOperator": "Operator",
    "SourceType": "SourceType", "Usage": "ListUsage", "Transformation": "Transformation",
    "NullPolicy": "NullPolicy", "InclusionRule": "InclusionRule",
    "ScenarioFamily": "ScenarioFamily", "ExistingECTDManager": "ExistingECTDManager",
    "SourceSystemCategory": "SourceSystemCategory", "SourceDatabaseType": "SourceDatabaseType",
    "PrimaryMigrationMethod": "PrimaryMigrationMethod", "TargetPlatform": "TargetPlatform",
    "SectionCode": "QuestionSection", "BaseConfidence": "Confidence",
}


def table_name(sheet: str) -> str:
    return "tblMvp" + "".join(part.title() for part in sheet.split("_", 1)[1].split("_"))


def json_field_map(cols: dict[str, list[str]]) -> list[dict]:
    """Create the initial map seed; normal builds read the maintained JSON file."""
    rows = []
    for sheet, root in JSON_ROOTS.items():
        source_key = "DerivationRuleId+ConditionGroup+ConditionSequence" if sheet == "03_Scenario_Derivation_Rules" else "ListCode+Code" if sheet == "22_Value_Lists" else cols[sheet][0]
        for column in cols[sheet]:
            if column in ("Notes", "ExampleValue", "ExampleInput", "ExampleOutput"):
                continue
            prop = column[:1].lower() + column[1:]
            dtype = "Boolean" if column in ("IsActive", "RuntimeExport", "SupportsMixedScope", "FallbackScenario", "StoreDetail") or column.startswith(("Is", "Supports", "Can", "Produces")) else "Integer" if column in ("Priority", "DisplaySequence", "ConditionSequence", "SortOrder", "Sequence") else "String"
            if sheet == "06_Requirement_Catalogue":
                nested = {"ModuleId": "owner", "OwningComponent": "owner", "PhaseScope": "scope", "ApplicabilityBasis": "scope", "ScenarioId": "scope", "ImplementationDisposition": "implementation", "ImplementationSheet": "implementation", "EngineCapability": "implementation", "SourceId": "source", "SourceSection": "source", "RequirementBasis": "source"}
                path = root + "." + nested[column] if column in nested else root
                prop = {"OwningComponent": "component", "ImplementationDisposition": "disposition", "ImplementationSheet": "sheet", "RequirementBasis": "basis"}.get(column, prop)
            else:
                path = root
            rows.append(dict(MappingId="JFM-" + sheet[:2] + "-" + re.sub(r"(?<!^)(?=[A-Z])", "-", column).upper(),
                SourceSheet=sheet, SourceTable=table_name(sheet), SourceColumn=column,
                SourceRecordKey=source_key, JSONPath=path, JSONProperty=prop, DataType=dtype,
                Transformation="GroupConditions" if column == "ConditionGroup" else "ToBoolean" if dtype == "Boolean" else "ToNumber" if dtype == "Integer" else "Direct",
                NullPolicy="Error" if column == cols[sheet][0] else "Omit", InclusionRule="ActiveAndApplicable",
                SortKey=source_key, SchemaVersion="0.1.0-mvp", ExampleInput="", ExampleOutput="",
                Notes="Column-level lineage. Nested projection is described by JSONPath."))
    return rows


def load_json_field_map() -> list[dict]:
    payload = json.loads(FIELD_MAP_SOURCE.read_text(encoding="utf-8"))
    if payload.get("schemaVersion") != "0.1.0-mvp" or not isinstance(payload.get("mappings"), list):
        raise ValueError("Maintained JSON field map has an unsupported schema or missing mappings")
    return payload["mappings"]


def _source_record_id(sheet: str, row: dict) -> str:
    if sheet == "03_Scenario_Derivation_Rules":
        return f"{row['DerivationRuleId']}:{row['ConditionGroup']}:{row['ConditionSequence']}"
    if sheet == "22_Value_Lists":
        return f"{row['ListCode']}:{row['Code']}"
    return str(row.get(columns()[sheet][0], ""))


def master_rows(rows: dict[str, list[dict]], selected: str) -> list[dict]:
    selected_maps = [m for m in rows["05_Scenario_Module_Map"] if m["ScenarioId"] == selected]
    module_apps = defaultdict(list)
    for m in selected_maps:
        module_apps[m["ModuleId"]].append(m["Applicability"])
    refs = {r["InputContextField"] for r in rows["03_Scenario_Derivation_Rules"]}
    refs.update(m["ActivationContextField"] for m in selected_maps if m["ActivationContextField"])
    result = []
    for sheet in SHEETS[1:24]:
        for row in rows[sheet]:
            record_type = {"01_Migration_Scenarios": "Scenario", "02_Scenario_Questionnaire": "Question",
                "03_Scenario_Derivation_Rules": "Rule", "04_Assessment_Modules": "Module",
                "05_Scenario_Module_Map": "ModuleMapping", "06_Requirement_Catalogue": "Requirement",
                "07_Fields_Evidence": "Field", "13_Size_Volume_Metrics": "Metric",
                "22_Value_Lists": "ValueList", "23_Source_References": "Source"}.get(sheet, "Rule")
            status, reason = "Excluded", "Outside the selected scenario or not approved for runtime export."
            phase = row.get("Phase", "All")
            module = row.get("ModuleId", "")
            app = ""
            if sheet == "01_Migration_Scenarios":
                if row["ScenarioId"] == selected:
                    status, reason = "Included", "Selected base scenario."
            elif sheet == "05_Scenario_Module_Map":
                if row["ScenarioId"] == selected:
                    app = row["Applicability"]
                    status = {"Required": "Included", "Conditional": "Conditional", "Optional": "Optional", "NotApplicable": "Excluded"}[app]
                    reason = row["BusinessReason"]
                else:
                    reason = "Mapping belongs to another scenario."
            elif sheet == "04_Assessment_Modules":
                if any(a != "NotApplicable" for a in module_apps[row["ModuleId"]]):
                    status, reason = "Included", "Referenced by a selected scenario-module mapping."
            elif sheet == "06_Requirement_Catalogue":
                if row["ImplementationDisposition"] == "Deferred":
                    status, reason = "Deferred", "Runtime implementation and rule review are pending."
                elif row["RuntimeExport"]:
                    status, reason = "Included", "Approved runtime requirement with JSON projection."
                else:
                    reason = "Authoring or validation obligation; not exported."
            elif sheet == "07_Fields_Evidence":
                if row["FieldCode"] in refs:
                    status, reason = "Included", "Referenced by derivation or selected activation logic."
            elif sheet == "13_Size_Volume_Metrics":
                if any(a != "NotApplicable" for a in module_apps[row["ModuleId"]]):
                    status, reason = "Included", "Approved PreSales measure used by an applicable module."
            elif sheet in ("02_Scenario_Questionnaire", "03_Scenario_Derivation_Rules", "22_Value_Lists", "23_Source_References"):
                status, reason = "Included", "Reusable approved definition or referenced source."
            elif row.get("IsActive") and (row.get("ScenarioId") in (selected, "ALL", None, "")):
                status, reason = "Included", "Active record in selected scenario scope."
            result.append(dict(SelectedScenarioId=selected, Phase=phase, ModuleId=module,
                ModuleApplicability=app, RequirementId=row.get("RequirementId", ""),
                RequirementTitle=row.get("RequirementTitle", ""),
                RequirementStatement=row.get("RequirementStatement", ""),
                SourceSheet=sheet, SourceRecordId=_source_record_id(sheet, row),
                RecordType=record_type, InclusionStatus=status, InclusionReason=reason,
                JSONPath=JSON_ROOTS[sheet] if status in ("Included", "Conditional", "Optional") else "",
                ReferencedBy="05_Scenario_Module_Map" if sheet == "04_Assessment_Modules" else "",
                ValidationStatus="Warning" if status == "Deferred" else "Valid",
                EngineCapability=row.get("EngineCapability", row.get("PrimaryEngineCapabilityGroup", ""))))
    return result


def preview_rows(rows: dict[str, list[dict]], master: list[dict], selected: str, issues: list[Issue]) -> list[dict]:
    scenario = next(r for r in rows["01_Migration_Scenarios"] if r["ScenarioId"] == selected)
    status = "Blocked" if any(i.blocking for i in issues) else "Eligible"
    active = [r for r in master if r["InclusionStatus"] in ("Included", "Conditional", "Optional")]
    counts = Counter(r["RecordType"] for r in active)
    sections = [
        ("configuration", 1, {"scenarioId": selected, "mappingVersion": "0.1.0", "schemaVersion": "0.1.0-mvp", "runtimeJsonEligible": status == "Eligible"}, "00_Home"),
        ("scenario", 1, {"id": selected, "code": scenario["ScenarioCode"], "name": scenario["ScenarioName"]}, "01_Migration_Scenarios"),
        ("modules[]", 45, {"selectedScenarioModuleRecords": 45, "notApplicableExplicit": True}, "05_Scenario_Module_Map"),
        ("questionnaire.questions[]", counts["Question"], {"firstQuestionId": rows["02_Scenario_Questionnaire"][0]["QuestionId"]}, "02_Scenario_Questionnaire"),
        ("questionnaire.derivationRules[]", len({r["DerivationRuleId"] for r in rows["03_Scenario_Derivation_Rules"]}), {"conditions": len(rows["03_Scenario_Derivation_Rules"])}, "03_Scenario_Derivation_Rules"),
        ("catalogues.fields[]", counts["Field"], {"fieldCodes": "See 07_Fields_Evidence"}, "07_Fields_Evidence"),
        ("requirements[]", sum(1 for r in rows["06_Requirement_Catalogue"] if r["RuntimeExport"]), {"draftLedgerRows": len(rows["06_Requirement_Catalogue"]), "exportBlocked": status == "Blocked"}, "06_Requirement_Catalogue"),
        ("rules.*[]", 0, {"approvedDerivationRulesAreInQuestionnaire": True, "detailedRuleSheetsPending": True}, "09_Dossier_Sequence_ID"),
        ("rules.sizeVolumeMetrics[]", counts["Metric"], {"phase": "PreSales", "basis": "Approximate planning measures"}, "13_Size_Volume_Metrics"),
        ("findings[]", counts["Finding"], {}, "18_Findings"),
        ("recommendations[]", counts["Recommendation"], {}, "19_Recommendations_Actions"),
        ("valueLists", counts["ValueList"], {"controlledCodeRows": counts["ValueList"]}, "22_Value_Lists"),
        ("sources[]", counts["Source"], {"sourceId": SOURCE_ID}, "23_Source_References"),
    ]
    return [dict(Section=section, ObjectCount=count, ValidationStatus=status,
        Preview=json.dumps(info, ensure_ascii=False, separators=(",", ":")), SourceSheet=source)
        for section, count, info, source in sections]


def issue_rows(issues: list[Issue], selected: str) -> list[dict]:
    return [dict(ValidationId=f"VAL-{i:04d}", Severity=issue.severity,
        ControlCode=issue.code, SheetName=issue.sheet, RecordId=issue.record,
        ColumnName=issue.column, Message=issue.message, WhyItMatters=issue.why,
        CorrectiveAction=issue.correction, Blocking=issue.blocking,
        SelectedScenarioId=selected) for i, issue in enumerate(issues, 1)]


def _coerce(column: str, value):
    if value == "":
        return None
    if column in ("SourceDate", "VerifiedOn") and isinstance(value, str):
        return date.fromisoformat(value)
    return value


def _width(header: str) -> int:
    if header == "Preview":
        return 72
    if header in ("Section", "SourceSheet", "SourceRecordId", "JSONPath"):
        return 38
    if header in ("ControlCode", "ReasonCode"):
        return 44
    if header in ("RequirementStatement", "AcceptanceCriterion", "BusinessPurpose", "BusinessReason", "AssessmentBoundary", "Description", "QuestionText", "InclusionReason", "Message", "CorrectiveAction", "WhyItMatters", "TechnicalImpact", "ReasonTemplate", "Guidance", "VerificationGuidance", "Rationale", "SourceSection"):
        return 52
    if header.endswith(("Id", "Code")) or header == "Reference":
        return 27
    return min(37, max(17, len(header) + 3))


def _add_table(ws, headers: list[str], records: list[dict], generated: bool) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "60758D" if generated else "245F8D"
    ws["A1"] = ws.title.replace("_", " · ", 1)
    ws["A1"].font = Font(name="Aptos", size=15, bold=True, color=NAVY)
    ws["A2"] = "← Home"
    ws["A2"].hyperlink = "#'00_Home'!A1"
    ws["A2"].font = Font(name="Aptos", size=10, color=BLUE, underline="single")
    ws["B2"] = "GENERATED · rerun builder" if generated else "MAINTAINED · review before runtime export"
    ws["B2"].font = Font(name="Aptos", size=10, color=MUTED, italic=True)
    for i, header in enumerate(headers, 1):
        cell = ws.cell(4, i, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = _width(header)
    ws.row_dimensions[4].height = 34
    material = records if records else [{}]
    for rnum, record in enumerate(material, 5):
        lines = 1
        for cnum, header in enumerate(headers, 1):
            value = _coerce(header, record.get(header, ""))
            cell = ws.cell(rnum, cnum, value)
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            wrap = header in ("BusinessDescription", "QuestionText", "BusinessMeaning", "BusinessReason", "RequirementStatement", "AcceptanceCriterion", "AssessmentBoundary", "Description", "Message", "CorrectiveAction", "WhyItMatters", "Preview", "InclusionReason", "Guidance", "VerificationGuidance", "ReasonTemplate", "Rationale", "TechnicalImpact")
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)
            if wrap and value is not None:
                lines = max(lines, 1 + len(str(value)) // max(12, _width(header) - 8))
            if isinstance(value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
            if header in ("RequirementId", "ScenarioId", "ModuleId", "FieldCode", "SourceRecordId"):
                cell.font = Font(name="Aptos", size=10, bold=True, color=BLUE)
            if header in ("RequirementStatus", "InclusionStatus", "ValidationStatus") and value in ("Deferred", "Warning", "Blocked"):
                cell.fill = PatternFill("solid", fgColor=AMBER)
        ws.row_dimensions[rnum].height = min(82, max(30, 12 + 15 * lines)) if records else 22
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4+len(material)}"
    ref = ws.auto_filter.ref
    table = Table(displayName=table_name(ws.title), ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    ws.freeze_panes = "E5" if ws.title == "06_Requirement_Catalogue" else "C5" if ws.title in ("05_Scenario_Module_Map", "24_Final_Config_Master") else "A5"
    ws.sheet_view.zoomScale = 85


def _name_ranges(wb, rows: dict[str, list[dict]], cols: dict[str, list[str]]) -> None:
    groups = defaultdict(list)
    for n, row in enumerate(rows["22_Value_Lists"], 5):
        groups[row["ListCode"]].append(n)
    for code, positions in groups.items():
        assert positions == list(range(positions[0], positions[-1] + 1))
        safe = re.sub(r"[^A-Za-z0-9_]", "_", code)
        wb.defined_names.add(DefinedName("VL_" + safe, attr_text=f"'22_Value_Lists'!$B${positions[0]}:$B${positions[-1]}"))
    for name, sheet, key in (("ScenarioCodes", "01_Migration_Scenarios", "ScenarioId"), ("ModuleCodes", "04_Assessment_Modules", "ModuleId"), ("QuestionCodes", "02_Scenario_Questionnaire", "QuestionId"), ("FieldCodes", "07_Fields_Evidence", "FieldCode"), ("SourceCodes", "23_Source_References", "SourceId"), ("RequirementCodes", "06_Requirement_Catalogue", "RequirementId")):
        c = cols[sheet].index(key) + 1
        letter = get_column_letter(c)
        wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${letter}$5:${letter}${4+len(rows[sheet])}"))
    wb.defined_names.add(DefinedName("SelectedScenarioId", attr_text="'00_Home'!$D$5"))


def _validations(wb, rows: dict[str, list[dict]], cols: dict[str, list[str]]) -> None:
    available = {r["ListCode"] for r in rows["22_Value_Lists"]}
    for sheet in SHEETS[1:]:
        ws = wb[sheet]
        if sheet in GENERATED:
            continue
        for idx, header in enumerate(cols[sheet], 1):
            list_name = CONTROLLED_COLUMNS.get(header)
            if sheet == "23_Source_References" and header == "VerificationStatus":
                list_name = "SourceVerificationStatus"
            if sheet == "25_JSON_Field_Map" and header == "DataType":
                list_name = "JSONDataType"
            if header == "ScenarioId" and sheet in ("13_Size_Volume_Metrics", "14_Source_DB_Archive_DMS", "15_RAG_Severity", "16_Confidence", "17_Effort_Drivers", "19_Recommendations_Actions", "20_PreMigration_Readiness", "21_PostMigration_Reconciliation"):
                formula = "=VL_ScenarioScope"
            elif header in ("ScenarioId", "CandidateScenarioId") and sheet != "01_Migration_Scenarios":
                formula = "=ScenarioCodes"
            elif header == "ModuleId" and sheet != "04_Assessment_Modules":
                formula = "=ModuleCodes"
            elif header in ("OriginQuestionId", "FollowUpQuestionId", "ParentQuestionId"):
                formula = "=QuestionCodes"
            elif header in ("FieldCode", "InputContextField", "ActivationContextField", "SourceFieldCode", "TargetFieldCode") and sheet != "07_Fields_Evidence":
                formula = "=FieldCodes"
            elif header == "SourceId" and sheet != "23_Source_References":
                formula = "=SourceCodes"
            elif header == "RequirementId" and sheet != "06_Requirement_Catalogue":
                formula = "=RequirementCodes"
            elif list_name in available:
                formula = "=VL_" + re.sub(r"[^A-Za-z0-9_]", "_", list_name)
            elif header == "IsActive" or header in ("RuntimeExport", "StoreDetail", "Blocking"):
                formula = '"TRUE,FALSE"'
            else:
                continue
            # Apply well beyond seeded content so new maintained rows inherit controls.
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            dv.error = "Choose an approved code from the controlled list."
            dv.errorTitle = "Invalid controlled value"
            dv.showErrorMessage = True
            dv.errorStyle = "stop"
            ws.add_data_validation(dv)
            letter = get_column_letter(idx)
            dv.add(f"{letter}5:{letter}{max(1000, 20+len(rows[sheet]))}")


def _home(wb, selected: str, issues: list[Issue], preview: list[dict]) -> None:
    ws = wb["00_Home"]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 37
    ws.column_dimensions["C"].width = 5
    ws.column_dimensions["D"].width = 26
    ws.column_dimensions["E"].width = 54
    ws["B2"] = "eMAS  /  MAPPING WORKBOOK"
    ws["B2"].font = Font(name="Aptos", size=18, bold=True, color=NAVY)
    ws["B3"] = "MVP Implementation Baseline v0.1"
    ws["B3"].font = Font(name="Aptos", size=13, color=BLUE)
    ws["B4"] = "Status: Draft for workbook content review"
    ws["B4"].font = Font(name="Aptos", size=11, bold=True, color="9A6200")
    ws["B5"] = "SelectedScenarioId"
    ws["D5"] = selected
    ws["D5"].fill = PatternFill("solid", fgColor=AMBER)
    ws["D5"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
    dv = DataValidation(type="list", formula1="=ScenarioCodes")
    dv.showErrorMessage = True
    dv.error = "Select an approved MS-01 through MS-08 scenario."
    ws.add_data_validation(dv)
    dv.add(ws["D5"])
    ws["B6"] = "Runtime JSON eligibility"
    ws["D6"] = "Blocked" if any(i.blocking for i in issues) else "Eligible"
    ws["D6"].font = Font(name="Aptos", size=11, bold=True, color="9C2F2F" if any(i.blocking for i in issues) else "216B44")
    ws["B7"] = "Validation"
    ws["D7"] = f"{sum(i.blocking for i in issues)} blocking · {sum(not i.blocking for i in issues)} warnings"
    ws["B8"] = "Preview freshness"
    ws["D8"] = f'=IF(D5="{selected}","Current","Stale — rerun builder")'
    ws["B10"] = "WORKFLOW"
    ws["B10"].font = Font(name="Aptos", size=11, bold=True, color=NAVY)
    ws["B11"] = "Edit maintained tables → regenerate workbook → inspect validation and JSON preview."
    ws["B12"] = "The PowerShell runtime consumes approved scenario JSON; it never reads this workbook."
    ws["B13"] = "No customer answers, executable code, VBA, or regulatory rules invented for review scaffolds."
    ws["B15"] = "NAVIGATION"
    ws["B15"].font = Font(name="Aptos", size=11, bold=True, color=NAVY)
    for i, name in enumerate(SHEETS[1:], 16):
        cell = ws.cell(i, 2, name)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = Font(name="Aptos", size=10, color=BLUE, underline="single")
        ws.cell(i, 4, "Generated" if name in GENERATED else "Maintained")
    ws["E5"] = "Scenario changes in Excel mark generated views stale; rerun with --scenario to rebuild them."
    ws["E6"] = "No active Runtime JSON is emitted while blocking validation remains."
    ws["E8"] = f"Preview: {preview[2]['ObjectCount']} module records, {preview[3]['ObjectCount']} questions, {preview[4]['ObjectCount']} derivation rules."
    for cell in (ws["E5"], ws["E6"], ws["E8"]):
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Aptos", size=10, color=MUTED)
    for key, definition in (("Scenario", "One approved base migration route; qualifiers stay separate."),
        ("Module", "Reusable bounded assessment capability mapped by scenario and phase."),
        ("Evidence", "Observed or supplied project input; answers are not stored here."),
        ("Finding", "Stable observation/result, separate from its action."),
        ("RAG", "Risk colour; independent of confidence and assessment status."),
        ("Confidence", "Evidence strength and coverage, not severity."),
        ("JSON", "Scenario-specific machine-readable configuration for the runtime.")):
        row = ws.max_row + 1
        ws.cell(row, 4, key)
        ws.cell(row, 5, definition)
    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 90
    ws.print_area = "B2:E49"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def _normalize_zip(path: Path) -> None:
    # openpyxl writes current ZIP entry timestamps. Normalize the container so
    # unchanged source content produces byte-identical workbook builds.
    with zipfile.ZipFile(path, "r") as source:
        entries = [(info.filename, source.read(info.filename), info.compress_type) for info in source.infolist()]
    entries = [(name, re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-09-13T12:00:00Z\g<2>", payload) if name == "docProps/core.xml" else payload, compression)
               for name, payload, compression in entries]
    fd, temp = tempfile.mkstemp(prefix="emas-mvp-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        with zipfile.ZipFile(temp, "w") as target:
            for name, payload, compression in entries:
                info = zipfile.ZipInfo(name, (2026, 9, 13, 12, 0, 0))
                info.compress_type = compression
                info.external_attr = 0o644 << 16
                target.writestr(info, payload)
        os.replace(temp, path)
        path.chmod(0o644)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def build(selected: str = "MS-01", output: Path = OUTPUT) -> dict:
    cols = columns()
    rows = source_rows()
    rows["25_JSON_Field_Map"] = load_json_field_map()
    issues = validate_model(rows, selected)
    master = master_rows(rows, selected)
    preview = preview_rows(rows, master, selected, issues)
    rows["24_Final_Config_Master"] = master
    rows["26_JSON_Preview"] = preview
    rows["27_Validation_Results"] = issue_rows(issues, selected)
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "eMAS"
    wb.properties.title = "eMAS Mapping Workbook MVP Implementation Baseline v0.1"
    wb.properties.description = "Draft for workbook content review; macro-free source configuration."
    wb.properties.created = datetime(2026, 9, 13, 12, 0, 0)
    wb.properties.modified = datetime(2026, 9, 13, 12, 0, 0)
    for sheet in SHEETS:
        ws = wb.create_sheet(sheet)
        if sheet != "00_Home":
            _add_table(ws, cols[sheet], rows[sheet], sheet in GENERATED)
    _name_ranges(wb, rows, cols)
    _validations(wb, rows, cols)
    _home(wb, selected, issues, preview)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    _normalize_zip(output)
    return {"path": str(output), "sheets": len(SHEETS), "scenarios": len(rows["01_Migration_Scenarios"]),
        "modules": len(rows["04_Assessment_Modules"]), "mappings": len(rows["05_Scenario_Module_Map"]),
        "requirements": len(rows["06_Requirement_Catalogue"]), "blocking": sum(i.blocking for i in issues),
        "warnings": sum(not i.blocking for i in issues)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--scenario", default="MS-01", help="Selected MS-01 through MS-08 scenario")
    parser.add_argument("--output", type=Path, default=OUTPUT, help="Override workbook output path for testing")
    args = parser.parse_args()
    print(json.dumps(build(args.scenario, args.output), indent=2))


if __name__ == "__main__":
    main()
